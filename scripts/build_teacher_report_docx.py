from __future__ import annotations

import argparse
import os
import re
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urlparse

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_HTML = Path(
    r"C:\Users\18453\AppData\Local\Temp\claude"
    r"\D--OceanEdge-Group-Marine-Debris-Detection-main"
    r"\fbe6405e-99a7-4d03-b938-ad34ed4c429d"
    r"\scratchpad\report.html"
)
DEFAULT_OUT = ROOT / "docs" / "paper" / "整体汇报_温金杰_2026-07-09.docx"

FONT = "Microsoft YaHei"
LATIN_FONT = "Calibri"
BLUE = RGBColor(0, 0, 0)
MID_BLUE = RGBColor(32, 32, 32)
DARK_BLUE = RGBColor(64, 64, 64)
GRAY = RGBColor(89, 89, 89)
LIGHT_GRAY = "F2F4F7"
NOTE_FILL = "F7F7F7"
KEY_FILL = "F2F2F2"
WHITE = "FFFFFF"
FIGURE_OVERRIDES = {
    "fig1_pipeline.png": ROOT / "docs" / "paper" / "figures" / "report_fig3_1_pipeline_paper.png",
    "fig2_modules.png": ROOT / "docs" / "paper" / "figures" / "report_fig4_1_cbam_bifpn_paper.png",
}


class Node:
    def __init__(self, tag: str, attrs: dict[str, str] | None = None) -> None:
        self.tag = tag
        self.attrs = attrs or {}
        self.children: list[Node | str] = []


class SimpleHTMLTree(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.stack = [self.root]

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        node = Node(tag.lower(), {k.lower(): v or "" for k, v in attrs})
        self.stack[-1].children.append(node)
        if tag.lower() not in {"br", "img", "hr", "meta", "link"}:
            self.stack.append(node)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        for idx in range(len(self.stack) - 1, 0, -1):
            if self.stack[idx].tag == tag:
                del self.stack[idx:]
                return

    def handle_data(self, data: str) -> None:
        if data:
            self.stack[-1].children.append(data)


def iter_nodes(node: Node, tag: str) -> list[Node]:
    found = []
    for child in node.children:
        if isinstance(child, Node):
            if child.tag == tag:
                found.append(child)
            found.extend(iter_nodes(child, tag))
    return found


def text_of(node: Node | str) -> str:
    if isinstance(node, str):
        return node
    if node.tag == "br":
        return "\n"
    return "".join(text_of(child) for child in node.children)


def clean_text(text: str) -> str:
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return text.strip()


def set_run_font(run, size: float | None = None, color: RGBColor | None = None, bold=None, italic=None):
    run.font.name = LATIN_FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph_border_bottom(paragraph, color: str = "D9D9D9", size: str = "8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx >= len(row.cells):
                continue
            row.cells[idx].width = Inches(width)
            tc_w = row.cells[idx]._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                row.cells[idx]._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_inline(paragraph, node: Node | str, *, inherited_bold=False, inherited_italic=False, size=11, color=None):
    if isinstance(node, str):
        parts = node.split("\n")
        for idx, part in enumerate(parts):
            if idx:
                paragraph.add_run().add_break()
            if part:
                run = paragraph.add_run(part)
                set_run_font(run, size=size, color=color, bold=inherited_bold, italic=inherited_italic)
        return

    bold = inherited_bold or node.tag in {"b", "strong"}
    italic = inherited_italic or node.tag in {"i", "em"}
    if node.tag == "br":
        paragraph.add_run().add_break()
        return
    if node.tag == "a":
        label = clean_text(text_of(node)) or node.attrs.get("href", "")
        run = paragraph.add_run(label)
        set_run_font(run, size=size, color=RGBColor(5, 99, 193), bold=bold, italic=italic)
        run.underline = True
        return
    for child in node.children:
        add_inline(paragraph, child, inherited_bold=bold, inherited_italic=italic, size=size, color=color)


def add_text_paragraph(doc: Document, node: Node, *, style=None, align=None, size=11, color=None) -> None:
    text = clean_text(text_of(node))
    if not text:
        return
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    for child in node.children:
        add_inline(p, child, size=size, color=color)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        paragraph_border_bottom(p)
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(
        run,
        size={1: 16, 2: 13, 3: 12}.get(level, 11),
        color={1: BLUE, 2: MID_BLUE, 3: DARK_BLUE}.get(level, DARK_BLUE),
        bold=True,
    )


def add_list(doc: Document, node: Node, ordered=False, level=0) -> None:
    idx = 1
    for child in node.children:
        if not isinstance(child, Node) or child.tag != "li":
            continue
        prefix = f"{idx}. " if ordered else "• "
        idx += 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25 + level * 0.2)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        marker = p.add_run(prefix)
        set_run_font(marker, size=10.5, color=BLUE if level == 0 else None, bold=not ordered)
        inline_children = []
        nested = []
        for item in child.children:
            if isinstance(item, Node) and item.tag in {"ul", "ol"}:
                nested.append(item)
            else:
                inline_children.append(item)
        for item in inline_children:
            add_inline(p, item, size=10.5)
        for nested_list in nested:
            add_list(doc, nested_list, ordered=nested_list.tag == "ol", level=level + 1)


def add_callout(doc: Document, node: Node, fill: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_widths(table, [6.5])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    for child in node.children:
        if isinstance(child, Node) and child.tag in {"ol", "ul"}:
            # Lists inside a callout render more reliably as short callout paragraphs.
            for li in [x for x in child.children if isinstance(x, Node) and x.tag == "li"]:
                p = cell.add_paragraph()
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.08
                run = p.add_run("• ")
                set_run_font(run, size=9.5, color=BLUE, bold=True)
                for item in li.children:
                    add_inline(p, item, size=9.5)
        else:
            text = clean_text(text_of(child))
            if text:
                p = cell.add_paragraph()
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.08
                add_inline(p, child, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def table_cell_text(cell_node: Node) -> str:
    return clean_text(text_of(cell_node))


def widths_for(cols: int) -> list[float]:
    if cols == 1:
        return [6.5]
    if cols == 2:
        return [1.75, 4.75]
    if cols == 3:
        return [1.65, 2.2, 2.65]
    if cols == 4:
        return [1.2, 1.75, 1.25, 2.3]
    if cols == 5:
        return [1.05, 1.2, 1.2, 1.2, 1.85]
    return [6.5 / cols] * cols


def add_table(doc: Document, node: Node) -> None:
    rows = []
    header_flags = []
    for tr in [x for x in node.children if isinstance(x, Node) and x.tag == "tr"]:
        cells = [x for x in tr.children if isinstance(x, Node) and x.tag in {"td", "th"}]
        if not cells:
            continue
        rows.append(cells)
        header_flags.append(any(c.tag == "th" for c in cells))
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    widths = widths_for(cols)
    set_table_widths(table, widths)
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            if c_idx >= len(row):
                continue
            cell_node = row[c_idx]
            if header_flags[r_idx]:
                shade_cell(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            text = table_cell_text(cell_node)
            align_center = c_idx > 0 and len(text) <= 24
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align_center else WD_ALIGN_PARAGRAPH.LEFT
            for item in cell_node.children:
                add_inline(
                    p,
                    item,
                    size=8.2 if cols >= 4 else 8.7,
                    inherited_bold=header_flags[r_idx],
                )
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def path_from_file_url(src: str) -> Path:
    parsed = urlparse(src)
    if parsed.scheme == "file":
        raw = unquote(parsed.path)
        if re.match(r"^/[A-Za-z]:/", raw):
            raw = raw[1:]
        return Path(raw.replace("/", os.sep))
    return Path(src)


def add_image(doc: Document, node: Node) -> None:
    src = node.attrs.get("src", "")
    if not src:
        return
    path = path_from_file_url(src)
    override = FIGURE_OVERRIDES.get(path.name)
    if override and override.exists():
        path = override
    if not path.exists():
        p = doc.add_paragraph()
        run = p.add_run(f"[图片缺失：{path}]")
        set_run_font(run, size=9, color=RGBColor(156, 0, 6), italic=True)
        return
    requested = float(node.attrs.get("width") or 620)
    width_inches = min(6.35, max(4.2, requested / 96.0))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = LATIN_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for level, size, color in [(1, 16, BLUE), (2, 13, MID_BLUE), (3, 12, DARK_BLUE)]:
        style = styles[f"Heading {level}"]
        style.font.name = LATIN_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("海洋垃圾检测项目整体汇报")
    set_run_font(run, size=9, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("温金杰 | 2026-07-09")
    set_run_font(run, size=9, color=GRAY)
    return doc


def add_static_toc(doc: Document, headings: list[tuple[int, str]]) -> None:
    add_heading(doc, "目录", 1)
    for level, title in headings:
        if level > 2:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25 * (level - 1))
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(title)
        set_run_font(run, size=10.5 if level == 1 else 10, color=BLUE if level == 1 else None, bold=level == 1)
    doc.add_page_break()


def build_docx(html_path: Path, out_path: Path) -> None:
    parser = SimpleHTMLTree()
    parser.feed(html_path.read_text(encoding="utf-8"))
    body = iter_nodes(parser.root, "body")[0]
    headings = []
    for child in body.children:
        if isinstance(child, Node) and child.tag in {"h2", "h3"}:
            headings.append((1 if child.tag == "h2" else 2, clean_text(text_of(child))))

    doc = setup_document()

    for child in body.children:
        if isinstance(child, str):
            continue
        if child.tag == "h1":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(clean_text(text_of(child)))
            set_run_font(run, size=22, color=BLUE, bold=True)
        elif child.tag == "h2":
            heading_text = clean_text(text_of(child))
            if heading_text.startswith("5.") or heading_text.startswith("附录 A") or heading_text.startswith("附录 B"):
                doc.add_page_break()
            add_heading(doc, heading_text, 1)
        elif child.tag == "h3":
            add_heading(doc, clean_text(text_of(child)), 2)
        elif child.tag == "h4":
            add_heading(doc, clean_text(text_of(child)), 3)
        elif child.tag == "p":
            text = clean_text(text_of(child))
            if text == "TOC_PLACEHOLDER":
                add_static_toc(doc, headings)
                continue
            klass = child.attrs.get("class", "")
            if "meta" in klass:
                add_text_paragraph(doc, child, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, color=RGBColor(64, 64, 64))
            elif "caption" in klass:
                add_text_paragraph(doc, child, align=WD_ALIGN_PARAGRAPH.CENTER, size=9, color=GRAY)
                if doc.paragraphs:
                    doc.paragraphs[-1].paragraph_format.keep_with_next = False
            else:
                add_text_paragraph(doc, child)
        elif child.tag == "div":
            klass = child.attrs.get("class", "")
            if "note" in klass:
                add_callout(doc, child, NOTE_FILL)
            elif "keybox" in klass:
                add_callout(doc, child, KEY_FILL)
            else:
                for item in child.children:
                    if isinstance(item, Node):
                        if item.tag == "p":
                            add_text_paragraph(doc, item)
                        elif item.tag in {"ul", "ol"}:
                            add_list(doc, item, ordered=item.tag == "ol")
        elif child.tag in {"ul", "ol"}:
            add_list(doc, child, ordered=child.tag == "ol")
        elif child.tag == "table":
            add_table(doc, child)
        elif child.tag == "img":
            add_image(doc, child)
        elif child.tag == "blockquote":
            add_callout(doc, child, WHITE)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def main() -> None:
    ap = argparse.ArgumentParser(description="Build the teacher-facing report DOCX from Claude's HTML draft.")
    ap.add_argument("--html", type=Path, default=DEFAULT_HTML)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()
    build_docx(args.html, args.out)
    print(args.out)


if __name__ == "__main__":
    main()
